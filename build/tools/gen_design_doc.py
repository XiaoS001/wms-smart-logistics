# -*- coding: utf-8 -*-
"""生成《概要设计说明书》DOCX。

格式遵循学校《本科毕业论文(设计)撰写规范》：
A4、上 30mm、下/左/右 25mm、页眉 20mm、页脚 15mm；
一级标题黑体小二、二级标题黑体四号、三级标题黑体小四；
正文宋体小四、1.25 倍行距、首行缩进 2 字符；表格三线表。
"""
import io
import os
import re
import time

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = r"D:\DeepSeekXM\XM01"
OUT_DIR = os.path.join(ROOT, "02-设计文档", "02-概要设计")
SRC = os.path.join(ROOT, "03-源码", "backend", "src", "main", "java", "com", "lijunjie", "wms")
CTRL = os.path.join(SRC, "controller")
SONG, HEI, TNR = "宋体", "黑体", "Times New Roman"


# --------------------------------------------------------------------------
# 排版
# --------------------------------------------------------------------------
def set_run_font(run, cn_font=SONG, size_pt=12, bold=False, ascii_font=TNR):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), cn_font)


def set_para(p, indent_chars=2, line_spacing=1.25, before_lines=0.0, after_lines=0.0, align=None):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pPr = p._element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if indent_chars:
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:firstLineChars"), str(int(indent_chars * 100)))
        ind.set(qn("w:firstLine"), str(int(indent_chars * 240)))
    elif ind is not None:
        pPr.remove(ind)
    pf.space_before = Pt(before_lines * 15)
    pf.space_after = Pt(after_lines * 15)
    if align is not None:
        pf.alignment = align
    return p


def H(doc, text, level):
    size, font = {1: (18, HEI), 2: (14, HEI), 3: (12, HEI)}[level]
    p = doc.add_paragraph()
    set_para(p, indent_chars=0, before_lines=0.5, after_lines=0.5,
             align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT)
    set_run_font(p.add_run(text), cn_font=font, size_pt=size, bold=True)
    return p


def B(doc, text, indent=2):
    p = doc.add_paragraph()
    set_para(p, indent_chars=indent, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    set_run_font(p.add_run(text))
    return p


def CODE(doc, text):
    p = doc.add_paragraph()
    set_para(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.LEFT, after_lines=0.2)
    set_run_font(p.add_run(text), cn_font=SONG, size_pt=10, ascii_font="Consolas")
    return p


def CAP(doc, text, size=10.5):
    p = doc.add_paragraph()
    set_para(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run_font(p.add_run(text), size_pt=size)
    return p


def tbl_borders(table, header_rows=1):
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    b = OxmlElement("w:tblBorders")
    for tag, sz in (("top", "8"), ("bottom", "8")):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), sz); e.set(qn("w:color"), "000000")
        b.append(e)
    for tag in ("left", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "none"); e.set(qn("w:sz"), "0")
        b.append(e)
    tblPr.append(b)
    for r in range(header_rows):
        for cell in table.rows[r].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(el)
            tb = OxmlElement("w:tcBorders")
            e = OxmlElement("w:bottom")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000")
            tb.append(e)
            tcPr.append(tb)


def T(doc, caption, headers, rows, widths=None, size=9):
    if caption:
        CAP(doc, caption)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ""
        p = c.paragraphs[0]; set_para(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run_font(p.add_run(h), cn_font=HEI, size_pt=size)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            p = c.paragraphs[0]
            set_para(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            set_run_font(p.add_run(str(v)), size_pt=size)
    tbl_borders(t, 1)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    sp = doc.add_paragraph(); set_para(sp, indent_chars=0)
    return t


# --------------------------------------------------------------------------
# 从源码统计真实规模，避免文档与实现脱节
# --------------------------------------------------------------------------
def count_java(pkg):
    d = os.path.join(SRC, pkg)
    if not os.path.isdir(d):
        return 0, 0
    files = [f for f in os.listdir(d) if f.endswith(".java")]
    lines = 0
    for f in files:
        with io.open(os.path.join(d, f), "r", encoding="utf-8") as fh:
            lines += len(fh.readlines())
    return len(files), lines


LAYERS = [
    ("controller", "接口层", "接收 HTTP 请求、参数校验、组织响应"),
    ("service", "业务接口层", "定义业务能力契约"),
    ("service\\impl", "业务实现层", "业务规则、事务边界、算法调用"),
    ("algorithm", "算法内核层", "货位分配与拣货路径策略，不依赖 Web 与数据库"),
    ("mapper", "数据访问层", "MyBatis-Plus 映射与自定义 SQL"),
    ("entity", "实体层", "与数据库表一一对应"),
    ("dto", "传输对象层", "接口出入参，与实体解耦"),
    ("common", "公共层", "统一响应、异常处理、工具类"),
]
layer_rows = []
for pkg, name, desc in LAYERS:
    n, lines = count_java(pkg)
    layer_rows.append([name, pkg.replace("\\", "/"), str(n), str(lines), desc])

algo_dir = os.path.join(SRC, "algorithm", "strategy")
slotting = sorted(f for f in os.listdir(algo_dir) if f.endswith("SlottingStrategy.java"))
picking = sorted(f for f in os.listdir(algo_dir) if f.endswith("PickingStrategy.java"))

ctrl_files = sorted(f for f in os.listdir(CTRL) if f.endswith(".java"))
MAPPING_RE = re.compile(r'@(Get|Post|Put|Delete)Mapping\b')
ctrl_rows = []
total_ep = 0
for f in ctrl_files:
    with io.open(os.path.join(CTRL, f), "r", encoding="utf-8") as fh:
        text = fh.read()
    n = len(MAPPING_RE.findall(text))
    total_ep += n
    ctrl_rows.append([f[:-5], str(n)])

# --------------------------------------------------------------------------
doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin, s.bottom_margin = Mm(30), Mm(25)
    s.left_margin, s.right_margin = Mm(25), Mm(25)
    s.header_distance, s.footer_distance = Mm(20), Mm(15)
st = doc.styles["Normal"]
st.font.name = TNR
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

p = doc.add_paragraph(); set_para(p, indent_chars=0, before_lines=1, after_lines=0.5, align=WD_ALIGN_PARAGRAPH.CENTER)
set_run_font(p.add_run("面向智慧物流的仓储管理系统"), cn_font=HEI, size_pt=16, bold=True)
p = doc.add_paragraph(); set_para(p, indent_chars=0, after_lines=1, align=WD_ALIGN_PARAGRAPH.CENTER)
set_run_font(p.add_run("概要设计说明书"), cn_font=HEI, size_pt=16, bold=True)

T(doc, "表 0.1　文档信息", ["项目", "内容"],
  [["文档名称", "概要设计说明书"], ["版本", "v1.0"],
   ["编制日期", time.strftime("%Y 年 %m 月 %d 日")],
   ["对应论文", "第 5 章 系统设计与实现"],
   ["上游文档", "《系统需求规格说明书》"],
   ["下游文档", "《数据库设计说明书》《算法设计说明书》《系统接口文档》"],
   ["设计依据", "毕业设计任务书 + 需求规格说明书 + 学校论文撰写规范"],
   ["实现规模", "%d 个控制器、%d 个接口、%d 张数据表" %
    (len(ctrl_files), total_ep, 18)]],
  widths=[3.0, 12.0])

# ============================== 第一章 ==============================
H(doc, "第一章　引言", 1)

H(doc, "1.1　编写目的", 2)
B(doc, "本文档依据《系统需求规格说明书》，给出面向智慧物流的仓储管理系统的概要设计，"
       "明确系统的总体结构、模块划分、核心组件职责、关键流程与部署方案，"
       "作为详细设计、编码实现与论文撰写的共同依据。")
B(doc, "本文档不描述每个类与方法的实现细节，这些内容由《数据库设计说明书》、"
       "《算法设计说明书》与《系统接口文档》分别承担。")

H(doc, "1.2　设计原则", 2)
B(doc, "（1）分层解耦：接口、业务、算法、数据访问四类职责严格分离，层间单向依赖，"
       "禁止接口层直接访问数据访问层；")
B(doc, "（2）算法独立：算法内核不依赖 Web 框架与数据库组件，只依赖 Java 标准库，"
       "因此可脱离系统单独运行与测试，也便于与 Python 原型交叉验证；")
B(doc, "（3）约定优于配置：统一响应体与统一异常处理，使各模块的接口行为一致；")
B(doc, "（4）数据一致：涉及多表变更的业务操作（入库确认、出库确认、盘点调整、方案应用）"
       "均置于同一事务中，失败整体回滚；")
B(doc, "（5）可复现：算法与实验均使用固定随机种子，结果可重复；")
B(doc, "（6）外部化配置：数据库口令与令牌密钥不进入版本库，通过环境变量注入。")

H(doc, "1.3　术语约定", 2)
B(doc, "本文档沿用《系统需求规格说明书》的术语。为便于阅读，再次说明三个核心概念：")
T(doc, "表 1.1　核心概念", ["概念", "含义"],
  [["货位与货位编号", "货位是仓库中的最小存储单元；算法内部使用从 0 开始的下标，"
    "按“到出入口距离升序、距离相同按主键升序”编号，该顺序即算法与数据库之间的映射规则"],
   ["目标函数", "加权拣货行走距离 F = Σ(商品出库频次 × 该商品所在货位到出入口的距离)，"
    "单位为“米·次”，是货位分配效果的统一度量口径"],
   ["拣货路径", "从出入口出发遍历订单内全部拣货点再返回出入口的行走顺序，"
    "其行走距离按曼哈顿度量计算"]],
  widths=[3.4, 11.6])

# ============================== 第二章 ==============================
H(doc, "第二章　总体设计", 1)

H(doc, "2.1　系统架构", 2)
B(doc, "系统采用前后端分离的 B/S 架构，后端按分层模式组织。整体结构自下而上分为"
       "数据层、数据访问层、算法内核层、业务实现层、接口层与表现层：")
CODE(doc, "表现层        Vue 3 前端（管理界面、库位平面图、路径可视化）")
CODE(doc, "                 ↓ HTTP / JSON")
CODE(doc, "接口层        Controller（参数校验、统一响应）")
CODE(doc, "                 ↓")
CODE(doc, "业务实现层    Service（业务规则、事务边界）  ←→  算法内核层（策略求解）")
CODE(doc, "                 ↓")
CODE(doc, "数据访问层    Mapper（MyBatis-Plus + 自定义 SQL）")
CODE(doc, "                 ↓")
CODE(doc, "数据层        MySQL 8.0（18 张表）")
B(doc, "其中算法内核层同时被业务实现层与测试代码调用：业务实现层在入库推荐货位、"
       "出库规划路径、执行货位优化时调用它；测试代码用 Python 原型生成的黄金数据"
       "对它做交叉验证。这种安排使算法既不脱离业务，又能独立验证。")

H(doc, "2.2　分层职责与实现规模", 2)
B(doc, "各层的职责与代码规模如下表。代码规模由源码自动统计，可随实现同步更新。")
T(doc, "表 2.1　分层职责与代码规模", ["层次", "包路径", "文件数", "代码行数", "主要职责"],
  layer_rows, widths=[2.4, 2.2, 1.4, 1.6, 7.4])

H(doc, "2.3　模块划分", 2)
B(doc, "系统划分为 9 个业务模块与 1 个算法模块。模块之间通过业务数据关联，"
       "不直接调用彼此的内部方法，依赖关系尽量单向：")
T(doc, "表 2.2　模块划分", ["序号", "模块", "职责", "对外接口数"],
  [["1", "用户与权限管理", "登录认证、用户维护、角色与权限分配", "3"],
   ["2", "仓库管理", "仓库基本信息与布局参数维护", "7"],
   ["3", "库位管理", "库区与货位维护、货位批量生成、近端货位检索", "2"],
   ["4", "商品管理", "商品与分类维护、出库频次累计、ABC 分类计算", "8"],
   ["5", "入库管理", "入库单创建、货位推荐、入库确认", "6"],
   ["6", "出库管理", "出库单创建、就近货位分配、拣货路径规划、出库确认", "8"],
   ["7", "库存管理", "库存查询、缺货与临期预警、盘点与差异调整", "8"],
   ["8", "货位优化", "优化方案预览、保存、应用与效果统计", "10"],
   ["9", "报表统计", "作业量趋势、库存分布、算法效果统计", "7"],
   ["10", "系统与算法支撑", "健康检查、算法求解与货位推荐", "2 + 3"]],
  widths=[1.2, 3.0, 8.6, 2.2])
B(doc, "说明：第 8 与第 10 项中部分接口共用一个控制器（货位分配算法接口同时服务于"
       "货位优化模块与入库模块的货位推荐），上表按业务归属分别列出。")

H(doc, "2.4　技术选型", 2)
T(doc, "表 2.3　技术选型与理由", ["层次", "技术", "版本", "选型理由"],
  [["开发语言", "Java", "JDK 1.8", "课程与任务书指定；生态成熟，企业应用广泛"],
   ["应用框架", "Spring Boot", "2.7.18", "与 JDK 8 兼容性最好的稳定版本；自动装配降低配置量"],
   ["持久层", "MyBatis-Plus", "3.5.7", "在 MyBatis 基础上提供通用 CRUD 与条件构造器，"
    "同时保留手写 SQL 的能力以应对复杂统计查询"],
   ["连接池", "Druid", "1.2.22", "内置监控与防 SQL 注入；初始连接 5、最大 50，支撑并发要求"],
   ["数据库", "MySQL", "8.0.27", "开源关系型数据库；InnoDB 支持事务与行级锁"],
   ["令牌", "jjwt", "0.9.1", "JWT 无状态认证，避免服务端维护会话"],
   ["工具库", "Hutool", "5.8.25", "简化加密、日期与字符串处理"],
   ["接口文档", "springdoc-openapi", "1.7.0", "基于注解生成在线接口说明，与代码同源"],
   ["算法原型", "Python + NumPy", "3.12", "便于快速实现与调参；作为 Java 实现的验证基准"],
   ["前端", "Vue 3", "—", "组件化开发，适合管理类界面"]],
  widths=[2.2, 3.0, 2.0, 8.0])

H(doc, "2.5　运行环境与部署架构", 2)
B(doc, "系统采用单机部署，适用于中小型仓库场景。部署结构如下：")
CODE(doc, "浏览器（Vue 3 单页应用）")
CODE(doc, "      │  HTTP :80 / :5173")
CODE(doc, "      ▼")
CODE(doc, "Web 服务器 / 前端静态资源")
CODE(doc, "      │  反向代理 /api/**")
CODE(doc, "      ▼")
CODE(doc, "Spring Boot 应用（内嵌 Tomcat，:8080）")
CODE(doc, "      │  JDBC（Druid 连接池）")
CODE(doc, "      ▼")
CODE(doc, "MySQL 8.0（:3306，库名 wms）")
B(doc, "应用以可执行 JAR 形式交付，启动命令为 java -jar wms-1.0.0.jar；"
       "数据库连接账号与令牌密钥通过环境变量注入，不写入代码库。")

# ============================== 第三章 ==============================
H(doc, "第三章　核心组件设计", 1)

H(doc, "3.1　统一响应体", 2)
B(doc, "所有接口共用同一响应结构，使前端可以用一套逻辑处理全部返回值：")
CODE(doc, '{ "code": 200, "message": "操作成功", "data": { ... }, "success": true, "timestamp": ... }')
B(doc, "code 为业务状态码，同时作为 HTTP 状态码返回，这样前端既能用 HTTP 状态做统一拦截"
       "（如 401 跳转登录页），也能直接读取 message 展示提示。")

H(doc, "3.2　统一异常处理", 2)
B(doc, "业务代码只抛出携带状态码与提示的 BusinessException，由全局异常处理器统一转换。"
       "处理器的关键设计是<b>按责任方区分日志级别</b>：客户端原因（参数不合法、请求体格式错误、"
       "方法不支持、内容类型不支持）记警告；服务端未预期异常记错误并保留堆栈。")
T(doc, "表 3.1　异常与状态码映射", ["异常类型", "HTTP", "业务码", "日志级别", "说明"],
  [["BusinessException", "按自带的码", "同左", "警告", "业务规则不满足，如库存不足、编码重复"],
   ["MethodArgumentNotValidException", "400", "400", "警告", "请求体字段校验失败"],
   ["BindException", "400", "400", "警告", "表单参数绑定失败"],
   ["HttpMessageNotReadableException", "400", "400", "警告", "请求体不是合法 JSON"],
   ["MissingServletRequestParameterException", "400", "400", "警告", "缺少必填参数"],
   ["HttpRequestMethodNotSupportedException", "405", "405", "警告", "请求方法不支持"],
   ["HttpMediaTypeNotSupportedException", "415", "415", "警告", "内容类型不支持"],
   ["其他 Exception", "500", "500", "错误", "未预期异常，不向前端暴露堆栈"]],
  widths=[5.4, 1.6, 1.4, 1.6, 5.0])

H(doc, "3.3　认证与权限", 2)
B(doc, "认证采用 JWT 无状态方案：登录接口校验口令后签发令牌，令牌内含用户名与有效期"
       "（默认 12 小时）。客户端在请求头以 Authorization: Bearer &lt;token&gt; 携带，"
       "服务端解析并校验签名与有效期。")
B(doc, "口令存储采用加盐迭代哈希（PBKDF2），同一口令两次加密结果不同，"
       "数据库中不保存明文。角色与权限采用“用户—角色—权限”三表关联模型，"
       "登录响应返回该用户的角色与权限编码列表，前端据此控制按钮显隐。")

H(doc, "3.4　算法内核", 2)
B(doc, "算法内核是本系统的设计重点。它只依赖 Java 标准库，不引入 Web 或数据库依赖，"
       "因此具备三个特性：可独立单元测试、可与 Python 原型交叉验证、可被多个业务模块复用。")
B(doc, "内核由三部分组成：仓库布局模型（负责坐标与曼哈顿距离计算）、算例模型"
       "（负责商品频次画像与 ABC 分类推导）、策略族（负责求解）。")
T(doc, "表 3.2　货位分配策略", ["策略编码", "名称", "思想", "复杂度"],
  [["random", "随机分配", "在全部货位中随机指定，作为效果对照基线", "O(n)"],
   ["frequency-greedy", "频次贪心", "商品按频次降序、货位按距离升序逐一配对", "O(n log n)"],
   ["abc-greedy", "ABC 分区贪心", "先按 ABC 分类划分近中远分区，再在类内按频次配对", "O(n log n)"],
   ["genetic", "遗传算法", "整数编码 + 锦标赛选择 + 均匀交叉 + 变异 + 精英保留", "O(G·P·n)"]],
  widths=[3.2, 2.4, 7.4, 2.0])
T(doc, "表 3.3　拣货路径策略", ["策略编码", "名称", "思想", "复杂度"],
  [["s-shape", "S 型穿越", "依次穿越含拣货点的巷道，方向交替", "O(n log n)"],
   ["largest-gap", "最大间隙", "在巷道内最大间隙处折返，减少无效穿越", "O(n log n)"],
   ["nearest-neighbor", "最近邻", "每次选择距当前位置最近的未访问拣货点", "O(n²)"],
   ["ant-colony", "蚁群算法", "信息素 + 启发式信息构造解，精英策略强化最优路径", "O(I·A·n²)"]],
  widths=[3.2, 2.4, 7.4, 2.0])
B(doc, "遗传算法的默认参数为种群 80、迭代 300、交叉率 0.9、变异率 0.15、"
       "锦标赛规模 3、精英数 2；蚁群算法默认蚂蚁 25、迭代 80、"
       "alpha = 1.0、beta = 3.0、挥发率 0.4、信息素强度 100、精英数 2。"
       "参数含义与取值依据见《算法设计说明书》。")

H(doc, "3.5　算法与业务的映射规则", 2)
B(doc, "算法在抽象的货位编号上求解，业务数据存放在数据库的货位表中，两者必须严格对应，"
       "否则求解结果无法落库与执行。本系统采用如下确定性映射规则：")
B(doc, "（1）货位编号：同一仓库内的货位按“到出入口距离升序、距离相同按主键升序”排列，"
       "编号 k 即该序列中的第 k 个货位；")
B(doc, "（2）商品编号：全部商品按主键升序排列，编号 i 对应第 i 个商品；")
B(doc, "（3）距离口径：算法内部用 aisleGap × x + rowGap × y 计算距离，"
       "与货位生成时写入数据库的 distance 字段使用同一公式，因此二者数值一致。")
B(doc, "这三条规则保证了“算法解 → 真实货位”的转换是确定且可复现的，"
       "也使论文中的实验数据与系统运行时的计算结果可以相互印证。")

H(doc, "3.6　控制器清单", 2)
T(doc, "表 3.4　控制器与接口数", ["控制器", "接口数"],
  [[r[0], r[1]] for r in ctrl_rows] + [["合计", str(total_ep)]],
  widths=[6.0, 3.0])

# ============================== 第四章 ==============================
H(doc, "第四章　核心流程设计", 1)

H(doc, "4.1　入库流程（算法推荐货位）", 2)
B(doc, "入库流程把货位分配算法嵌入业务操作，是系统由“记录工具”变为“决策支持工具”的第一个落点。")
CODE(doc, "仓管员                入库服务              算法内核              数据库")
CODE(doc, "  │  创建入库单        │                     │                    │")
CODE(doc, "  ├───────────────────>│                     │                    │")
CODE(doc, "  │                    │ 校验仓库与明细      │                    │")
CODE(doc, "  │                    ├─────────────────────────────────────────>│ 查询商品/货位")
CODE(doc, "  │                    │ 生成单号 IN+日期+流水号                  │")
CODE(doc, "  │                    ├─────────────────────────────────────────>│ 写入单据")
CODE(doc, "  │  useSlottingSuggest=true                 │                    │")
CODE(doc, "  │                    │ 构造算例（布局+频次）│                   │")
CODE(doc, "  │                    ├────────────────────>│                    │")
CODE(doc, "  │                    │                     │ 按频次降序 × 距离升序配对")
CODE(doc, "  │                    │<────────────────────┤ 返回目标值与分配方案")
CODE(doc, "  │                    │ 映射为真实货位并写入推荐理由             │")
CODE(doc, "  │                    ├─────────────────────────────────────────>│ 写入明细")
CODE(doc, "  │  返回单据与推荐货位 │                     │                    │")
CODE(doc, "  │<───────────────────┤                     │                    │")
B(doc, "关键设计：<b>人工指定优先于算法推荐</b>。明细行若已填写货位，算法不再覆盖；"
       "算法不可用（例如无空闲货位）时只记录警告，不阻塞单据创建，保证业务连续性。")
B(doc, "入库确认时在<b>同一事务</b>内完成四件事：写入或累加库存、把货位置为占用、"
       "记录实收数量、把单据置为已完成。任何一步失败整体回滚，避免出现"
       "“库存增加了但货位未占用”这类不一致状态。")

H(doc, "4.2　出库流程（就近分配 + 拣货路径规划）", 2)
B(doc, "出库流程是算法的第二个落点，分为分配与规划两个阶段：")
CODE(doc, "仓管员                出库服务              算法内核              数据库")
CODE(doc, "  │  创建出库单        │                     │                    │")
CODE(doc, "  ├───────────────────>│                     │                    │")
CODE(doc, "  │                    │ 查该商品全部库存    │                    │")
CODE(doc, "  │                    ├─────────────────────────────────────────>│ 查询库存")
CODE(doc, "  │                    │ 按货位到出入口距离升序依次扣减（就近出库）│")
CODE(doc, "  │                    │ 库存不足则返回具体差额                   │")
CODE(doc, "  │                    ├─────────────────────────────────────────>│ 写入明细")
CODE(doc, "  │  规划拣货路径       │                     │                    │")
CODE(doc, "  ├───────────────────>│ 汇总订单拣货点坐标  │                    │")
CODE(doc, "  │                    ├────────────────────>│                    │")
CODE(doc, "  │                    │                     │ 四种策略分别求解并对比")
CODE(doc, "  │                    │<────────────────────┤ 返回顺序与总距离")
CODE(doc, "  │                    │ 顺序写回明细 pick_seq，距离与算法写入单据 │")
CODE(doc, "  │                    ├─────────────────────────────────────────>│ 更新")
B(doc, "就近出库的意义：同一商品的库存可能分布在多个货位，优先从离出入口最近的货位出库，"
       "可以缩短该次拣货的行走距离，同时使近端货位更快周转。")
B(doc, "出库确认时在同一事务内扣减库存、释放已清空的货位、标记明细已拣，"
       "并<b>累加商品出库频次</b>——频次正是货位分配算法的输入，"
       "由此形成“业务运行 → 频次更新 → 算法优化”的数据闭环。")

H(doc, "4.3　货位优化方案流程", 2)
B(doc, "货位优化面向全局，处理“商品应该放在哪里”的问题，流程分为预览、保存、应用三步：")
CODE(doc, " 预览：求解优化方案 → 计算当前实际布局的目标值 → 对比得出降幅与搬迁清单")
CODE(doc, " 保存：把方案编号、策略、参数、前后目标值、降幅与分配明细写入方案表")
CODE(doc, " 应用：按方案把库存记录搬迁到目标货位，并在同一事务内重算货位占用状态")
B(doc, "“当前目标值”取自库存表的实际数据（商品实际所在货位 × 其出库频次），"
       "“优化后目标值”取自算法求解结果，两者口径一致，因此降幅具有可比性。")
B(doc, "应用阶段的实现要点：搬迁过程中<b>不逐条维护货位状态</b>。"
       "因为“商品 A 的原货位”可能正是“商品 B 的目标货位”，"
       "若在循环内立即释放原货位，会把另一个商品刚占用的货位错误置为空闲。"
       "正确做法是循环内只搬数据，循环结束后按最终库存统一重算该仓库的货位占用状态。")
B(doc, "此外，为避免同一商品的多个批次争抢同一个目标货位，应用时只搬迁数量最大的一条"
       "库存记录，其余批次保持原位，并在返回结果中明确提示，需要人工处理的批次不会被静默忽略。")

H(doc, "4.4　库存盘点流程", 2)
B(doc, "盘点用于校正账面库存与实物库存的差异，流程为：")
CODE(doc, "生成盘点单（以当前库存为账面数量）→ 录入实盘数量 → 计算差异 → 可选地调整库存")
B(doc, "调整库存时在同一事务内：按实盘数量更新或删除库存记录、同步货位占用状态、"
       "把盘点记录标记为已调整。已调整的盘点单不能重复提交或撤销，"
       "保证库存变动有据可查、不可追溯地篡改。")

H(doc, "4.5　货位批量生成流程", 2)
B(doc, "货位数据是算法的基础输入，系统提供按仓库布局批量生成的能力：")
CODE(doc, "读取仓库布局（巷道数 X、纵深数 Y、巷道间距 a、纵深间距 b）")
CODE(doc, "  ↓ 生成 X × Y 个货位，坐标 (x, y)")
CODE(doc, "计算 distance = a × x + b × y（与算法内核同一公式）")
CODE(doc, "  ↓ 按距离阈值划分 A/B/C 库区")
CODE(doc, "批量写入并按坐标去重（重复调用不会产生重复货位）")
B(doc, "生成过程具有幂等性：已存在的坐标默认跳过，需要重建时可显式指定覆盖。"
       "单次最多生成 200 × 200 个货位，超过则拒绝，防止误操作拖垮数据库。")

# ============================== 第五章 ==============================
H(doc, "第五章　数据库与部署设计", 1)

H(doc, "5.1　数据库总体设计", 2)
B(doc, "数据库共 18 张表，按用途分为四类：")
T(doc, "表 5.1　数据表分类", ["类别", "表名", "用途"],
  [["系统管理", "sys_user、sys_role、sys_permission、sys_user_role、sys_role_permission、sys_operation_log",
    "用户、角色、权限及关联关系与操作日志"],
   ["基础数据", "base_warehouse、base_zone、base_location、base_category、base_product",
    "仓库布局、库区、货位、商品分类与商品主数据"],
   ["业务单据", "biz_inbound_order、biz_inbound_item、biz_outbound_order、biz_outbound_item、biz_stocktake、biz_slotting_plan",
    "入库、出库、盘点单及其明细，以及货位优化方案"],
   ["库存", "inv_stock", "商品在货位上的实时库存"]],
  widths=[2.0, 6.4, 6.6])
B(doc, "表结构的字段定义、索引与约束详见《数据库设计说明书》。"
       "此处仅说明三条贯穿全库的设计约定：")
B(doc, "（1）逻辑删除：业务表统一使用 deleted 标记，查询自动过滤，避免误删导致历史单据丢失；")
B(doc, "（2）单据唯一：单号（入库单、出库单、盘点单、优化方案）均建立唯一约束，"
       "防止重复提交产生孪生单据；")
B(doc, "（3）明细唯一：单据明细的唯一约束按“单号 + 对象”组合，"
       "例如盘点明细为“单号 + 货位 + 商品”。若错误地对单号单独建唯一约束，"
       "一张单据就只能有一条明细，插入第二条即报重复键错误。")

H(doc, "5.2　部署设计", 2)
B(doc, "系统为单机部署，组件与端口如下：")
T(doc, "表 5.2　部署组件", ["组件", "端口", "说明"],
  [["MySQL 8.0", "3306", "数据库服务，字符集 utf8mb4，存储引擎 InnoDB"],
   ["Spring Boot 应用", "8080", "内嵌 Tomcat，提供全部 REST 接口"],
   ["前端静态资源", "80 或 5173", "开发态由 Vite 提供，生产态由 Nginx 托管"],
   ["接口文档（在线）", "8080", "springdoc-openapi 生成的 /swagger-ui 页面"]],
  widths=[3.6, 2.4, 9.0])
B(doc, "启动顺序：先启动数据库，再启动后端应用，最后访问前端。"
       "应用提供健康检查接口 /api/health，一次调用即可确认"
       "“应用已启动 + 数据库连通 + 表结构已导入”，可用于部署自检与演示前检查。")

H(doc, "5.3　安全设计", 2)
T(doc, "表 5.3　安全措施", ["措施", "实现方式"],
  [["口令保护", "PBKDF2 加盐迭代哈希存储，数据库中无明文口令"],
   ["身份认证", "JWT 无状态令牌，默认有效期 12 小时，支持环境变量配置密钥"],
   ["权限控制", "用户—角色—权限三级模型，登录响应返回权限编码供前端控制按钮"],
   ["敏感配置", "数据库口令与令牌密钥通过环境变量注入，不进入版本库"],
   ["信息不泄露", "登录失败不区分“账号不存在”与“口令错误”；用户查询接口不返回口令字段；"
    "异常响应不返回堆栈"],
   ["参数校验", "接口层校验必填项与取值范围，业务层校验业务规则，数据库层用约束兜底"]],
  widths=[3.0, 12.0])

H(doc, "5.4　性能设计", 2)
T(doc, "表 5.4　性能措施", ["措施", "说明"],
  [["连接池", "Druid 初始连接 5、最大 50，支撑并发访问"],
   ["索引", "货位距离、商品出库频次、单据状态与时间等高频查询字段建立索引；"
    "“按距离取近端空闲货位”这一算法高频查询可直接命中索引，无需全表排序"],
   ["批量写入", "货位批量生成采用分批插入，避免拼接超长 SQL"],
   ["限流保护", "单次货位生成上限 200 × 200；列表类接口设置返回条数上限"],
   ["算法降级", "拣货点较多时可选用计算更快的策略，保证响应时间"]],
  widths=[2.6, 12.4])

# ============================== 第六章 ==============================
H(doc, "第六章　设计约束与可扩展性", 1)

H(doc, "6.1　当前设计的约束", 2)
B(doc, "（1）服务单一仓库：系统按单仓库场景设计，暂不支持多仓库并行与跨仓调拨；")
B(doc, "（2）一货位一商品：同一货位在同一时刻只存放一个商品（允许多批次），"
       "这与算法“商品与货位一一配对”的建模一致；")
B(doc, "（3）规则网格布局：货位按巷道与纵深规则排列，行走距离按曼哈顿度量计算，"
       "未考虑不规则仓库与障碍物；")
B(doc, "（4）静态频次：出库频次由历史出库记录累计，算法按当前频次求解，"
       "不做需求预测。")

H(doc, "6.2　可扩展方向", 2)
B(doc, "（1）多仓库：在仓库相关表中补充分仓字段与数据权限，即可扩展为多仓场景；")
B(doc, "（2）需求预测：把静态频次替换为基于时间序列的预测频次，使优化具备前瞻性；")
B(doc, "（3）容量约束：当前模型假设一货位一商品且不设容量上限，"
       "可引入容量与重量约束，使问题更贴近实际（算法侧需相应扩展目标函数与约束处理）；")
B(doc, "（4）新策略：算法内核采用策略模式与工厂注册，新增分配或路径策略只需实现接口并在工厂登记，"
       "无需改动业务代码；")
B(doc, "（5）实时可视化：货位平面图与拣货路径已在接口层返回坐标序列，"
       "前端可直接绘制，后续可扩展为动效演示。")

doc.save(os.path.join(OUT_DIR, "概要设计说明书-v1-%s.docx" % time.strftime("%Y%m%d")))
out_path = os.path.join(OUT_DIR, "概要设计说明书-v1-%s.docx" % time.strftime("%Y%m%d"))
print("saved:", out_path)
print("size: %.1f KB" % (os.path.getsize(out_path) / 1024.0))
print("controllers=%d endpoints=%d" % (len(ctrl_files), total_ep))
